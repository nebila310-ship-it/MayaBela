#!/usr/bin/env python3
"""Build the LIA – Majo Bridge software licence and implementation agreement as an editable Word file.

Usage: python3 docs/proposals/build_lia_agreement_docx.py
Needs `pip install python-docx`.

Figures come from the one-time financial proposal MBT-LIA-FIN-2026-02 (49 modules, 16 October 2026).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "MaJo_Bridge_LIA_Agreement.docx"
LOGO = ROOT.parent.parent / "assets" / "branding" / "maya_brand.png"

BLUE = RGBColor(0x1A, 0x4F, 0x9C)
GREY = RGBColor(0x4A, 0x55, 0x68)

# ----- commercial figures (MBT-LIA-FIN-2026-02) -----
MODULES_ETB = 3_425_000
IMPLEMENTATION_ETB = 500_000
DISCOUNT_ETB = 50_000
SUBTOTAL_ETB = MODULES_ETB + IMPLEMENTATION_ETB - DISCOUNT_ETB  # 3,875,000
VAT_ETB = int(SUBTOTAL_ETB * 0.15)  # 581,250
TOTAL_ETB = SUBTOTAL_ETB + VAT_ETB  # 4,456,250
PAY1 = int(TOTAL_ETB * 0.40)  # 1,782,500
PAY2 = int(TOTAL_ETB * 0.40)  # 1,782,500
PAY3 = TOTAL_ETB - PAY1 - PAY2  # 891,250

WORDS = {
    TOTAL_ETB: "Four Million Four Hundred Fifty-Six Thousand Two Hundred Fifty",
    SUBTOTAL_ETB: "Three Million Eight Hundred Seventy-Five Thousand",
    VAT_ETB: "Five Hundred Eighty-One Thousand Two Hundred Fifty",
    PAY1: "One Million Seven Hundred Eighty-Two Thousand Five Hundred",
    PAY3: "Eight Hundred Ninety-One Thousand Two Hundred Fifty",
}


def etb(n: int) -> str:
    return f"{n:,} ETB"


def etb_words(n: int) -> str:
    return f"{etb(n)} ({WORDS[n]} Ethiopian Birr)"


MODULES = [
    ("Academic and assessment", [
        (1, "Weighted Markbook", "Premium", 95_000), (2, "LMS", "Premium", 95_000),
        (3, "Grade Workflow", "Premium", 95_000), (4, "Examinations & Grade Approvals", "Premium", 95_000),
        (5, "Exam bank & papers", "High", 75_000), (6, "Academic Management", "High", 75_000),
        (7, "Curriculum office", "High", 75_000), (8, "Homework", "High", 75_000),
        (9, "Lesson plans", "High", 75_000), (10, "Report Cards", "High", 75_000),
        (11, "Timetable", "High", 75_000), (12, "Analytics & exports", "Mid", 65_000),
        (13, "e-Books & Materials", "Mid", 65_000),
    ]),
    ("Student information and admissions", [
        (14, "Students (SIS)", "Premium", 95_000), (15, "Attendance", "Premium", 95_000),
        (16, "Admissions", "High", 75_000), (17, "Transfers & Promotion", "High", 75_000),
        (18, "At-risk & insights", "High", 75_000), (19, "Student support", "High", 75_000),
        (20, "Parents & Link Approvals", "High", 75_000), (21, "Add Student", "Mid", 65_000),
        (22, "Student Affairs & Discipline", "Mid", 65_000), (23, "Student programs", "Mid", 65_000),
        (24, "Alumni", "Standard", 50_000),
    ]),
    ("Procurement and store", [
        (25, "Inventory, Procurement & Store", "High", 75_000),
    ]),
    ("Safeguarding, quality and compliance", [
        (26, "Safeguarding", "Premium", 95_000), (27, "Reports & Analytics", "Premium", 95_000),
        (28, "Go-live & compliance", "Premium", 95_000), (29, "Audit Log", "High", 75_000),
        (30, "QA Findings & Plans", "Mid", 65_000),
    ]),
    ("Communication and school life", [
        (31, "Messages", "High", 75_000), (32, "Calendar", "Mid", 65_000),
        (33, "Announcements", "Mid", 65_000), (34, "Library", "Mid", 65_000),
        (35, "Events", "Mid", 65_000), (36, "Gallery", "Mid", 65_000),
        (37, "Maya Assistant", "Mid", 65_000),
    ]),
    ("Administration and access", [
        (38, "Role Permissions", "Premium", 95_000), (39, "Institution Management", "Mid", 65_000),
        (40, "School Management", "Mid", 65_000), (41, "System Health", "Standard", 50_000),
        (42, "Classroom Teachers", "Standard", 50_000), (43, "Add Teacher", "Standard", 50_000),
        (44, "Administration Staff Directory", "Standard", 50_000),
        (45, "Add Administration Staff", "Standard", 50_000),
        (46, "Student Portal Settings", "Lite", 35_000), (47, "Student Password Resets", "Lite", 35_000),
        (48, "Settings", "Lite", 35_000), (49, "Dashboard / Profile", "Lite", 35_000),
    ]),
]

assert sum(p for _, rows in MODULES for *_, p in rows) == MODULES_ETB
assert sum(len(rows) for _, rows in MODULES) == 49

DELIVERABLES = [
    ("D1", "Inception report (MBT-LIA-INC-2026-01)", "Scope, approach, plan, governance, risks, Client inputs", "M2", "LIA Procurement / project sponsor"),
    ("D2", "Software requirements and customization registers", "Requirements per office, customization per requirement, priorities, LIA owners (Annex E of the Inception Report)", "M2", "LIA focal person and office heads"),
    ("D3", "Hybrid environment", "School node installed on the LIA server; cloud node prepared; tunnel and sub-domain; one-login test; backup job; architecture sheet as built", "M3", "IT Admin"),
    ("D4", "Configured and customized LIA school", "School record, campuses, academic year, terms, grading schemes, approval chains, role map, branding, customizations from D2", "M3", "IT Admin and Principal"),
    ("D5", "Migrated data", "Students, parents and links, staff and roles, classes, subjects, allocations, timetable on the school node; import log and exception list", "M4", "Registrar"),
    ("D6", "Training delivered", "Sessions per the Training Modules (MBT-LIA-TRN-2026-01), attendance sheets, in-app manuals, train-the-trainer", "M5", "Principal"),
    ("D7", "Restore drill and security check", "Backup restored to a test copy and verified; security checklist of both nodes", "M5", "IT Admin"),
    ("D8", "Pilot report", "Two weeks of one section live: usage, issues found and fixed, parent-link results", "W7", "Section Director of the pilot section"),
    ("D9", "UAT report and acceptance certificate", "Scripted scenarios run by LIA officers, defect log, retest results, exit-criteria check", "M6", "LIA project sponsor"),
    ("D10", "Go-live certificate", "All sections live; first backup verified; support contact card in every office", "M7", "LIA project sponsor"),
    ("D11", "Hypercare closure and support plan", "Open issues with owners, support channels, escalation path, operations handbook for LIA IT", "M8", "IT Admin"),
]

MILESTONES = [
    ("M1", "Kick-off held; requirements workshops complete", "Fri 30 Oct 2026"),
    ("M2", "Requirements and customization registers signed; final Inception Report accepted", "Fri 6 Nov 2026"),
    ("M3", "Both nodes live and connected; LIA configured; admin and IT can sign in", "Fri 13 Nov 2026"),
    ("M4", "Roster and staff list imported to the school node and signed", "Fri 20 Nov 2026"),
    ("M5", "Training complete; restore drill passed; pilot section live", "Fri 4 Dec 2026"),
    ("M6", "User-acceptance testing passed; delivery accepted", "Fri 18 Dec 2026"),
    ("M7", "Go-live, all sections", "Mon 21 Dec 2026"),
    ("M8", "Hypercare closed; support period starts", "Fri 15 Jan 2027"),
]

SERVER_SPEC = [
    ("Server", "1 rack or tower server, 8-core CPU, 32 GB RAM, 2 × 1 TB SSD in mirror (RAID 1), dual gigabit network ports"),
    ("Operating system", "Ubuntu Server 22.04 LTS or 24.04 LTS with disk encryption, installed by the Service Provider with LIA IT present"),
    ("Power", "UPS of 1.5–2 kVA with at least 20 minutes of runtime for the server and switch"),
    ("Room", "Locked, ventilated or air-conditioned room; access list held by IT Admin"),
    ("Network", "Static LAN address; gigabit switch to staff rooms; Wi-Fi reaching classrooms; internet line of at least 20 Mbps up and down"),
    ("Domain", "One sub-domain under the Academy's domain pointing to the tunnel"),
    ("Backup disks", "2 × 2 TB external encrypted disks; one with the server, one kept in another building and swapped weekly"),
]


# ----- document helpers -----

def set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def apply_widths(table, widths_cm: list[float]) -> None:
    """Fix column widths so Word and LibreOffice both honour them."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        col.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


class Builder:
    def __init__(self) -> None:
        self.doc = Document()
        self.article_no = 0
        self.clause_no = 0
        self._styles()
        self._page()

    def _styles(self) -> None:
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, color in (("Heading 1", 14, BLUE), ("Heading 2", 12, BLUE), ("Heading 3", 11, GREY)):
            h = st[name]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            rpr = h.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            rfonts.set(qn("w:ascii"), "Calibri")
            rfonts.set(qn("w:hAnsi"), "Calibri")

    def _page(self) -> None:
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        hdr = sec.header.paragraphs[0]
        hdr.text = "Software Licence and Implementation Agreement · Lebawi International Academy and Majo Bridge Technologies and Events PLC"
        hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hdr.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = GREY
        ftr = sec.footer.paragraphs[0]
        ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(ftr, "Page ", "PAGE", " of ", "NUMPAGES", "   ·   Initials: LIA ______   MBT ______")

    @staticmethod
    def _field(par, *parts) -> None:
        for p in parts:
            if p in ("PAGE", "NUMPAGES"):
                for kind, text in (("begin", None), ("instr", p), ("separate", None), ("text", "1"), ("end", None)):
                    run = par.add_run()
                    run.font.size = Pt(8)
                    run.font.color.rgb = GREY
                    if kind == "instr":
                        it = OxmlElement("w:instrText")
                        it.set(qn("xml:space"), "preserve")
                        it.text = text
                        run._r.append(it)
                    elif kind == "text":
                        run.text = text
                    else:
                        fc = OxmlElement("w:fldChar")
                        fc.set(qn("w:fldCharType"), kind)
                        run._r.append(fc)
            else:
                r = par.add_run(p)
                r.font.size = Pt(8)
                r.font.color.rgb = GREY

    # --- text ---
    def p(self, text: str = "", bold: bool = False, italic: bool = False, align=None, size: float | None = None, color=None, after: float | None = None):
        par = self.doc.add_paragraph()
        if text:
            run = par.add_run(text)
            run.bold = bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
        if align is not None:
            par.alignment = align
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        return par

    def rich(self, parts: list[tuple[str, bool]], align=None):
        par = self.doc.add_paragraph()
        for text, bold in parts:
            par.add_run(text).bold = bold
        if align is not None:
            par.alignment = align
        return par

    def h1(self, text: str):
        return self.doc.add_heading(text, level=1)

    def h2(self, text: str):
        return self.doc.add_heading(text, level=2)

    def article(self, title: str):
        self.article_no += 1
        self.clause_no = 0
        return self.doc.add_heading(f"Article {self.article_no} — {title}", level=2)

    def clause(self, text: str, parts: list[tuple[str, bool]] | None = None):
        self.clause_no += 1
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.first_line_indent = Cm(-1.0)
        par.add_run(f"{self.article_no}.{self.clause_no}\t").bold = True
        if parts:
            for t, b in parts:
                par.add_run(t).bold = b
        else:
            par.add_run(text)
        par.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        return par

    def sub(self, items: list[str], letters: bool = True):
        for i, item in enumerate(items):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(2.0)
            par.paragraph_format.first_line_indent = Cm(-0.8)
            par.paragraph_format.space_after = Pt(3)
            marker = f"({chr(97 + i)})" if letters else "•"
            par.add_run(f"{marker}\t{item}")
            par.paragraph_format.tab_stops.add_tab_stop(Cm(2.0))

    def bullets(self, items: list[str]):
        for item in items:
            par = self.doc.add_paragraph(style="List Bullet")
            par.add_run(item)
            par.paragraph_format.space_after = Pt(2)

    def table(self, header: list[str], rows: list[list[str]], widths: list[float] | None = None, size: float = 9.5, bold_last: bool = False):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0]
        set_repeat_header(hdr)
        for i, h in enumerate(header):
            c = hdr.cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = BLUE
            set_cell_shading(c, "F4F7FB")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(size)
                if bold_last and ri == len(rows) - 1:
                    r.bold = True
        if widths:
            apply_widths(t, widths)
        for row in t.rows:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.line_spacing = 1.0
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self):
        self.doc.save(OUT)


# ----- content -----

def cover(b: Builder) -> None:
    if LOGO.exists():
        b.doc.add_picture(str(LOGO), width=Cm(3.2))
        b.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    b.p("SOFTWARE LICENCE AND IMPLEMENTATION AGREEMENT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=18, color=BLUE, after=2)
    b.p("for the Student Management System (SMS) Project", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, after=14)
    b.p("between", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    b.p("LEBAWI INTERNATIONAL ACADEMY", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, after=0)
    b.p("(hereinafter “the Client” or “LIA”)", align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, after=10)
    b.p("and", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    b.p("MAJO BRIDGE TECHNOLOGIES AND EVENTS PLC", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, after=0)
    b.p("(hereinafter “the Service Provider” or “Majo Bridge”)", align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, after=18)

    b.table(
        ["Item", "Detail"],
        [
            ["Agreement reference", "MBT-LIA-AGR-2026-01"],
            ["Subject", "One-time licence of 49 MaJo e-School Bridge modules with four login portals, customization, hybrid implementation, data migration, training, go-live, and one year of support for Lebawi International Academy"],
            ["Contract price", f"{etb_words(TOTAL_ETB)}, inclusive of VAT at 15%"],
            ["Payment", "40% on signing · 40% on delivery for user-acceptance testing · 20% on go-live"],
            ["Place of signing", "Addis Ababa, Federal Democratic Republic of Ethiopia"],
            ["Date of signing", "______ day of ______________ 2026 G.C. (______ ______________ 2019 E.C.)"],
            ["Annexes", "A Module list and prices · B Deliverables and milestones · C Payment schedule · D Client-provided infrastructure · E Inception Report MBT-LIA-INC-2026-01 · F Training Modules MBT-LIA-TRN-2026-01 · G Financial Proposal MBT-LIA-FIN-2026-02 · H Notice of Intent to Award · I Certificate of Recognition (form)"],
        ],
        widths=[4.2, 12.4],
        size=10,
    )
    b.p("Prepared in two (2) original counterparts in the English language, one for each Party. Underscored blanks are completed before signature.", italic=True, color=GREY, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    b.page_break()


def parties(b: Builder) -> None:
    b.h1("THIS AGREEMENT")
    b.p("This Software Licence and Implementation Agreement (the “Agreement”) is made and entered into in Addis Ababa, Ethiopia, on this ______ day of ______________ 2026 G.C. (______ ______________ 2019 E.C.),")
    b.p("BETWEEN", bold=True, after=2)
    b.rich([
        ("LEBAWI INTERNATIONAL ACADEMY, ", True),
        ("a private educational institution established and operating under the laws of the Federal Democratic Republic of Ethiopia, with its principal office at Addis Ababa, ______________ Sub-City, Woreda ______, House No. ______, Tel. ______________, Tax Identification Number (TIN) ______________, Business/Operating Licence No. ______________, represented for the purpose of this Agreement by ______________________________, its ______________________________ (hereinafter referred to as the “Client” or “LIA”), of the one part;", False),
    ])
    b.p("AND", bold=True, after=2)
    b.rich([
        ("MAJO BRIDGE TECHNOLOGIES AND EVENTS PLC, ", True),
        ("a private limited company incorporated and registered under the laws of the Federal Democratic Republic of Ethiopia, with its principal office at Addis Ababa, ______________ Sub-City, Woreda ______, House No. ______, Tel. +251 911 646 444 / +251 911 133 548, e-mail majobridgetech@gmail.com, Tax Identification Number (TIN) ______________, Commercial Registration No. ______________, Business Licence No. ______________, represented for the purpose of this Agreement by ______________________________, its General Manager (hereinafter referred to as the “Service Provider” or “Majo Bridge”), of the other part.", False),
    ])
    b.p("The Client and the Service Provider are hereinafter referred to individually as a “Party” and collectively as the “Parties”.")

    b.h2("PREAMBLE")
    b.p("WHEREAS the Client issued Terms of Reference for a Student Management System to replace spreadsheet grading, paper summaries, and informal messaging groups across its Elementary, Middle School, and High School sections;")
    b.p("WHEREAS the Service Provider submitted a Technical Proposal (v1.0 FINAL, 12 September 2026), an Addendum on Hybrid Hosting (15 September 2026), and a one-time Financial Proposal for forty-nine (49) modules, document MBT-LIA-FIN-2026-02 dated 16 October 2026, for its product MaJo e-School Bridge;")
    b.p("WHEREAS the Client, through its Procurement Department, issued a Notice of Intent to Award for the Student Management System (SMS) Project naming the Service Provider as the preferred service provider, subject to contract negotiation and agreement on the final scope and terms;")
    b.p("WHEREAS the Service Provider has tabled an Inception Report, document MBT-LIA-INC-2026-01 dated 27 October 2026, setting out the confirmed scope, the software requirements analysis, the hybrid implementation approach, the work plan, the key deliverables, and the governance of the project, together with Training Modules, document MBT-LIA-TRN-2026-01;")
    b.p("WHEREAS the Parties wish to record the terms and conditions on which the Service Provider shall licence, customize, implement, and support the System for the Client;")
    b.p("NOW, THEREFORE, in consideration of the mutual covenants herein contained, and in accordance with the Civil Code of the Empire of Ethiopia of 1960 (Proclamation No. 165/1960), in particular the provisions governing contracts in general and contracts for the performance of services, the Parties hereby agree as follows:", bold=False)


def articles(b: Builder) -> None:
    # 1 Definitions
    b.article("Definitions and Interpretation")
    b.clause("In this Agreement, unless the context otherwise requires, the following terms have the meanings given below:")
    b.sub([
        "“System” means the MaJo e-School Bridge software, comprising the forty-nine (49) Modules listed in Annex A, the four (4) login portals (Administration/ERP, Teacher, Parent, Student), the Customizations, and the related in-app manuals, as installed and configured for the Client.",
        "“Module” means a functional component of the System listed and priced in Annex A.",
        "“Customization” means a change to the default behaviour, forms, rules, or layout of the System recorded in the Customization Register (Deliverable D2) and agreed in writing by both Parties.",
        "“Inception Report” means document MBT-LIA-INC-2026-01, Version 1.0 of 27 October 2026, attached as Annex E, as finalised and accepted under Milestone M2, and any later revision signed by both Parties.",
        "“Hybrid Implementation” means the architecture described in Section 6 and Annex D of the Inception Report, in which the School Node runs on the Client's physical server and the Cloud Node runs on the Service Provider's self-hosted cloud server.",
        "“School Node” means the installation of the System's database, data API, file storage, and live-update services on the Client's physical server at the Client's premises, holding the master copy of student, academic, care, document, and message records.",
        "“Cloud Node” means the installation of the System's authentication, role catalogue, application delivery, notification, and system-health services on the Service Provider's self-hosted cloud server, holding no student records.",
        "“Client Data” means all data entered into, uploaded to, or generated by the System on behalf of the Client, including personal data of students, parents, and staff.",
        "“Deliverable” and “Milestone” mean the items listed in Annex B, as detailed in Sections 7 and 8 of the Inception Report.",
        "“Go-Live” means the date on which the System is in productive use by all sections of the Client and the Go-Live Certificate (Deliverable D10) is signed, targeted for Monday 21 December 2026.",
        "“User-Acceptance Testing” or “UAT” means the scripted scenarios in Section 13 of the Inception Report, run by the Client's officers on the Client's own data.",
        "“Support Period” means the period of twelve (12) months beginning on the Go-Live date during which support is provided free of charge under Article 12.",
        "“Contract Price” means the amount stated in Article 4.",
        "“Working Day” means Monday to Friday, excluding public holidays recognised in Ethiopia.",
    ])
    b.clause("Headings are for convenience only. Words in the singular include the plural and vice versa. References to Articles and Annexes are to Articles of and Annexes to this Agreement. The Gregorian calendar governs all dates unless the Ethiopian calendar is expressly stated.")

    # 2 Documents
    b.article("Contract Documents and Order of Precedence")
    b.clause("The following documents form an integral part of this Agreement and are incorporated by reference:")
    b.sub([
        "this Agreement, Articles 1 to 24;",
        "Annex A — Module list, categories, and unit prices;",
        "Annex B — Deliverables and milestones;",
        "Annex C — Payment schedule;",
        "Annex D — Infrastructure, premises, and inputs to be provided by the Client;",
        "Annex E — Inception Report MBT-LIA-INC-2026-01, including its Annexes A to E (module checklist, data templates, training timetable, hybrid architecture and server specification, requirements and customization registers);",
        "Annex F — Training Modules and Trainer's Guide MBT-LIA-TRN-2026-01;",
        "Annex G — Financial Proposal MBT-LIA-FIN-2026-02 of 16 October 2026;",
        "Annex H — Notice of Intent to Award issued by the Client's Procurement Department;",
        "Annex I — Form of Certificate of Recognition.",
    ])
    b.clause("In the event of conflict or inconsistency, the documents shall prevail in the order listed in Article 2.1, except that the Contract Price and the payment terms in Articles 4 and 5 prevail over any amount or payment term in Annexes E to H.")
    b.clause("The Technical Proposal and the Addendum on Hybrid Hosting are background documents. They describe the product and may be consulted for interpretation but do not add obligations beyond those in this Agreement and its Annexes.")

    # 3 Scope
    b.article("Scope of Work")
    b.clause("The Service Provider shall, for the Contract Price:")
    b.sub([
        "grant to the Client the licence described in Article 11 for all forty-nine (49) Modules in Annex A and the four (4) login portals;",
        "conduct the software requirements analysis with the Client's offices and produce the Requirements Register and the Customization Register (Deliverable D2);",
        "implement the System under the Hybrid Implementation: install, configure, and secure the School Node on the Client's physical server and the Cloud Node on the Service Provider's self-hosted cloud server, connect them, and demonstrate the one-login test (Deliverable D3);",
        "configure and customize the System for the Client: school record, campuses, academic year and terms, grading schemes, approval chains, role map, branding, and the Customizations in the Customization Register (Deliverable D4);",
        "migrate the Client's data from the completed templates to the School Node and deliver the import log and exception list (Deliverable D5);",
        "deliver the training sessions in Annex F and the in-app manuals (Deliverable D6);",
        "perform the backup restore drill and the security check of both nodes (Deliverable D7);",
        "run a two-week pilot in one section and report on it (Deliverable D8);",
        "support the Client's User-Acceptance Testing, fix defects, and obtain the acceptance certificate (Deliverable D9);",
        "bring all sections live and issue the Go-Live Certificate (Deliverable D10);",
        "provide hypercare for four (4) weeks after Go-Live, including on-site presence in the first week, and hand over the operations handbook (Deliverable D11); and",
        "provide support free of charge during the Support Period under Article 12.",
    ])
    b.clause("The following are outside the scope of this Agreement unless added by a written variation under Article 15: Finance, Payroll, Human Resource, Live GPS, Transport, Buses, Digital Operations, Campus Management, CCTV, Register Driver, and the Driver portal; hardware and network equipment; recurring third-party charges (SMS gateway, internet lines, domain renewals, payment gateways); re-typing of records not provided in the agreed templates; and support, hosting, or subscription after the Support Period.")
    b.clause("The Client may, at the kick-off meeting, confirm a smaller selection of not fewer than twenty-five (25) Modules. In that case the module component of the Contract Price is recalculated as the sum of the unit prices in Annex A for the Modules confirmed, plus the implementation fee, less the discount, plus VAT, in accordance with Annex G, and Annex A and Annex C are re-issued and initialled by both Parties before the first invoice.")

    # 4 Price
    b.article("Contract Price")
    b.clause("", [("The Contract Price for the full scope in Article 3 is ", False), (etb_words(TOTAL_ETB), True), (", inclusive of Value Added Tax at fifteen percent (15%), computed as follows:", False)])
    b.table(
        ["Component", "Amount (ETB)"],
        [
            ["Forty-nine (49) Modules at the published unit prices (Annex A)", f"{MODULES_ETB:,}"],
            ["Customization, implementation, and go-live (one-time)", f"{IMPLEMENTATION_ETB:,}"],
            ["Discount before VAT", f"({DISCOUNT_ETB:,})"],
            ["Sub-total, VAT exclusive", f"{SUBTOTAL_ETB:,}"],
            ["VAT at 15%", f"{VAT_ETB:,}"],
            ["Contract Price, VAT inclusive", f"{TOTAL_ETB:,}"],
        ],
        widths=[12.0, 4.6],
        size=10,
        bold_last=True,
    )
    b.clause("The Contract Price is a fixed, lump-sum price in Ethiopian Birr. It includes the licence, all Deliverables in Annex B, travel and subsistence of the Service Provider's staff within Addis Ababa, the Cloud Node hosting during the implementation and the Support Period, and one year of support. No escalation applies.")
    b.clause("The Contract Price does not include the items in Article 3.2 or the infrastructure in Annex D, which are procured by the Client at its own cost.")

    # 5 Payment
    b.article("Terms of Payment")
    b.clause("The Client shall pay the Contract Price in three (3) instalments as follows:")
    b.table(
        ["Instalment", "Share", "Amount (ETB)", "Due on", "Supporting document"],
        [
            ["First", "40%", f"{PAY1:,}", "Signing of this Agreement", "Signed Agreement; VAT invoice of the Service Provider"],
            ["Second", "40%", f"{PAY2:,}", "Delivery of the configured Modules for User-Acceptance Testing (UAT readiness notice issued after Milestone M5)", "UAT readiness notice; Deliverables D3 to D8 accepted or deemed accepted; VAT invoice"],
            ["Third", "20%", f"{PAY3:,}", "Go-Live (Milestone M7)", "Signed Go-Live Certificate (D10); VAT invoice"],
            ["Total", "100%", f"{TOTAL_ETB:,}", "", ""],
        ],
        widths=[2.4, 1.4, 2.6, 5.2, 5.0],
        size=9,
        bold_last=True,
    )
    b.clause("", [("The first instalment of ", False), (etb_words(PAY1), True), (" is due and payable upon signing of this Agreement, against the Service Provider's VAT invoice, and shall be paid within five (5) Working Days of signature. The Service Provider is not obliged to mobilise, install the Cloud Node, or begin configuration before the first instalment is received.", False)])
    b.clause("The second and third instalments shall be paid within ten (10) Working Days of receipt of the Service Provider's VAT invoice together with the supporting document in Article 5.1.")
    b.clause("Payments shall be made in Ethiopian Birr by bank transfer or cheque to the account of the Service Provider: Bank ______________________, Branch ______________________, Account No. ______________________, Account name Majo Bridge Technologies and Events PLC. Bank charges of the paying bank are borne by the Client.")
    b.clause("The Client shall withhold tax from each payment as required by the Federal Income Tax Proclamation No. 979/2016 and its regulations and shall deliver the official withholding tax receipt to the Service Provider within fifteen (15) days of payment. Amounts properly withheld and evidenced by such receipt are deemed paid to the Service Provider.")
    b.clause("If an undisputed instalment remains unpaid thirty (30) days after its due date, the Service Provider may, after seven (7) days' written notice, suspend the services until payment is received, and the dates in Annex B shall be extended by the period of suspension. Interest at the legal rate under the Civil Code shall accrue on the overdue amount.")
    b.clause("If the Client disputes any part of an invoice, it shall notify the Service Provider in writing within five (5) Working Days of receipt, stating the reasons, and shall pay the undisputed part on the due date.")

    # 6 Duration
    b.article("Effective Date, Duration, and Time Schedule")
    b.clause("This Agreement enters into force on the date of its signature by both Parties (the “Effective Date”) and remains in force until the end of the Support Period and the fulfilment of all obligations under it.")
    b.clause("The implementation shall follow the twelve-week work plan in Section 7 of the Inception Report, from the kick-off meeting of Tuesday 27 October 2026 to the end of hypercare on Friday 15 January 2027, with Go-Live targeted for Monday 21 December 2026. The Milestones in Annex B are binding dates subject to Articles 6.3 and 17.")
    b.clause("Milestone M3 and every later Milestone depend on the Client's physical server being installed in its room, powered, and on the network by Friday 6 November 2026 (Annex D) and on the Client's other inputs being delivered by the dates in Section 15 of the Inception Report. If an input is late, the Milestones that depend on it move by the same number of Working Days without penalty to the Service Provider, and the Service Provider shall not place Client Data on the Cloud Node to save time.")
    b.clause("The Parties confirm dates and rooms at the kick-off meeting and record them in the minutes. Changes to dates agreed in the steering meeting are recorded in writing and do not require an amendment to this Agreement unless they change the Contract Price.")

    # 7 Service Provider obligations
    b.article("Obligations of the Service Provider")
    b.clause("The Service Provider shall:")
    b.sub([
        "perform the services with the skill, care, and diligence expected of a professional software provider, using qualified personnel, and in accordance with the Inception Report;",
        "assign a project lead as the single point of contact, attend the weekly steering meeting, and issue a one-page progress report every second Friday;",
        "install the School Node software on the Client's server with the Client's IT Administrator present, and hand the server administrator account and the backup encryption key to the Client's IT Administrator;",
        "keep all Client Data on the School Node; the Cloud Node shall hold no student, academic, care, or document records, which the Client may verify under UAT scenario 16 of the Inception Report;",
        "configure nightly encrypted backups of the School Node to the Client's disks and perform the restore drill before Go-Live;",
        "access Client Data only through named maintenance accounts that appear in the audit log, and only for configuration, import, training, and support;",
        "deliver the training sessions in Annex F and keep the in-app manuals current;",
        "correct, at no cost, defects in the System reported during implementation, UAT, hypercare, and the Support Period;",
        "comply with the Client's reasonable site rules, child-protection rules, and the school calendar, including the no-change window before semester examinations; and",
        "maintain the Cloud Node, apply updates with notice in the project group, and update the School Node in a maintenance window agreed with the IT Administrator.",
    ])

    # 8 Client obligations
    b.article("Obligations of the Client")
    b.clause("The Client shall:")
    b.sub([
        "pay the Contract Price in accordance with Article 5;",
        "procure, install, and maintain at its own cost the physical server, UPS, server room, network, internet line, sub-domain, and backup disks specified in Annex D, by Friday 6 November 2026;",
        "appoint, at the kick-off meeting, a project sponsor, a project focal person available about two hours per Working Day, an IT Administrator available about one day per week, and the office counterparts listed in Section 9 of the Inception Report;",
        "complete and return the data templates in Annex B of the Inception Report, with one spelling per class name and correct dates of birth and parent telephone numbers, by Friday 13 November 2026;",
        "make the officers and teachers available for the requirements workshops, the training sessions in Annex F, the pilot, and UAT at the dates agreed;",
        "review each Deliverable within five (5) Working Days and give written comments or acceptance; silence after five (5) Working Days is treated as acceptance;",
        "operate the System after Go-Live: marks, attendance, parent-link approvals, announcements, first-line help through the trained helpers, server room, power, weekly disk swap, and the weekly IT checklist;",
        "hold the server administrator account and the backup encryption key securely and keep the access list of the server room;",
        "obtain any consents from parents and staff required under applicable law for the processing of their personal data in the System; and",
        "designate the School Node sub-domain and give the Service Provider the DNS access needed to point it to the tunnel.",
    ])

    # 9 Hybrid & data
    b.article("Hybrid Implementation, Data Residency, and Ownership of Data")
    b.clause("The System shall be implemented under the Hybrid Implementation. Student information, academic records, care and safeguarding records, documents and large files, and messages shall reside on the School Node at the Client's premises. Authentication, the role catalogue, application delivery, published announcements, calendar and events, notification delivery, and system health shall run on the Cloud Node.")
    b.clause("All Client Data is and remains the exclusive property of the Client. The Service Provider acquires no right in Client Data other than the limited right to process it for the purposes of this Agreement.")
    b.clause("Users sign in once at the Cloud Node; the School Node accepts sessions signed by the Cloud Node. The session-signing key is held by both Parties. Row-level security and role-based access apply on both nodes.")
    b.clause("Off-campus access to the School Node is provided through an outbound encrypted tunnel on the Client's sub-domain; no inbound port is opened on the Client's router. On campus, staff reach the School Node over the LAN; users already signed in continue to work during an internet outage, while new sign-ins wait for the line to return.")
    b.clause("On expiry or termination of this Agreement, the Client retains the School Node and all Client Data on it. The Service Provider shall, on request and at no cost within the Support Period, export the Client Data in a standard open format and delete any Client Data held temporarily by it, except as required by law.")

    # 10 Delivery, acceptance
    b.article("Delivery, Testing, and Acceptance")
    b.clause("Each Deliverable is submitted to the acceptor named in Annex B with a transmittal note. The acceptor reviews it within five (5) Working Days. Comments are answered in writing by the Service Provider; silence after five (5) Working Days is treated as acceptance.")
    b.clause("The second instalment becomes due when the Service Provider issues the UAT readiness notice confirming that the configured Modules, the migrated data, and the trained users are ready for User-Acceptance Testing, after Milestone M5.")
    b.clause("UAT is run by the Client's officers using the scenarios in Section 13 of the Inception Report on the Client's own data under the Client's control. Defects are logged with a severity. The Service Provider fixes defects and the Client retests.")
    b.clause("Acceptance and Go-Live criteria: no open Severity-1 or Severity-2 defect; at least ninety percent (90%) of scenarios passed first time and one hundred percent (100%) after retest; roster signed; officers trained; restore drill passed; scenario 16 (no student records on the Cloud Node) passed; support card in every office. On fulfilment, the Client signs the acceptance certificate (D9) and, on the agreed Go-Live date, the Go-Live Certificate (D10).")
    b.clause("Severity-3 and Severity-4 defects open at Go-Live are listed in the hypercare closure note with target dates and do not delay Go-Live or the third instalment.")
    b.clause("Use of the System in production by the Client for fifteen (15) consecutive Working Days without a written notice of non-acceptance constitutes acceptance.")

    # 11 Licence & IP
    b.article("Licence and Intellectual Property")
    b.clause("The Service Provider grants to the Client, from the Effective Date and subject to payment of the Contract Price, a perpetual, non-exclusive, non-transferable licence to install, access, and use the Modules in Annex A and the four login portals for the Client's own educational operations at its campuses in Ethiopia, for an unlimited number of the Client's staff, students, and parents.")
    b.clause("All intellectual property rights in the System, including its source code, object code, design, documentation, and any Customization, improvement, or derivative work created in the performance of this Agreement, are and remain the property of the Service Provider. Nothing in this Agreement transfers ownership of the System to the Client.")
    b.clause("The Client shall not sell, sub-license, lease, lend, or make the System available to any third party, reverse engineer or decompile it except as permitted by law, remove proprietary notices, or use it for any school or institution other than the Client without the Service Provider's written consent.")
    b.clause("The Service Provider warrants that it owns or lawfully controls the System and that the Client's use of it under this Agreement does not infringe the rights of any third party, and shall indemnify the Client against any claim to the contrary.")
    b.clause("The Service Provider may use the knowledge and generic techniques gained in the project for other clients, provided no Client Data or confidential information of the Client is used or disclosed.")

    # 12 Support & warranty
    b.article("Warranty, Hypercare, and Support Period")
    b.clause("The Service Provider warrants that, for the Support Period, the System shall perform materially in accordance with the Inception Report, the Requirements Register, and the Customization Register, and shall correct any non-conformity at no cost.")
    b.clause("Hypercare: for four (4) weeks after Go-Live (W9 to W12), the Service Provider shall be on site during the first week, check in daily, and respond the same day to every issue.")
    b.clause("Support Period: for twelve (12) months from Go-Live, support is provided free of charge through telephone and WhatsApp (+251 911 646 444, +251 911 133 548) and e-mail (majobridgetech@gmail.com, nabilmaya6464@gmail.com), Monday to Saturday, 09:00 to 18:00 East Africa Time, with the following response targets:")
    b.table(
        ["Severity", "Definition", "Response", "Target resolution or workaround"],
        [
            ["Sev-1", "School cannot sign in; a node is down; data loss risk", "Within 2 hours, any day", "Within 8 hours"],
            ["Sev-2", "A desk or portal unusable for a group of users; no workaround", "Same Working Day", "Within 2 Working Days"],
            ["Sev-3", "Function impaired; workaround exists", "Within 2 Working Days", "Next scheduled release"],
            ["Sev-4", "Cosmetic; question; enhancement request", "Within 5 Working Days", "Considered for a future release"],
        ],
        widths=[1.6, 6.4, 3.6, 5.0],
        size=9,
    )
    b.clause("Support covers logins, roles, the software of both nodes, the tunnel, the Android application, releases, and restore assistance. It does not cover the Client's hardware, power, internet line, or operational tasks listed in Article 8.1(g), nor training of new staff beyond the train-the-trainer scheme, which may be requested at the Service Provider's then-current rates.")
    b.clause("Support, hosting of the Cloud Node, and software updates after the Support Period are not included in the Contract Price. Not later than sixty (60) days before the end of the Support Period, the Service Provider shall submit a proposal for continued support and hosting, which the Client may accept or decline. If declined, the Client retains its licence and the School Node; the Service Provider may decommission the Cloud Node sixty (60) days after the end of the Support Period, after offering to migrate the authentication service to a server of the Client's choice at a price to be agreed.")

    # 13 Confidentiality & data protection
    b.article("Confidentiality and Protection of Personal Data")
    b.clause("Each Party shall keep confidential all information of the other Party marked as confidential or which by its nature is confidential, including Client Data, the Service Provider's source code and pricing, and the contents of this Agreement, and shall use it only for the purposes of this Agreement. This obligation survives for five (5) years after the end of this Agreement and, for personal data, indefinitely.")
    b.clause("The Service Provider shall process personal data in the System only on the documented instructions of the Client and in compliance with the Personal Data Protection Proclamation No. 1321/2024 of Ethiopia and other applicable law; shall implement the technical and organisational measures in Section 12 of the Inception Report (TLS encryption, disk and backup encryption, row-level security, audit logs, key custody); shall ensure its personnel are bound by confidentiality; shall notify the Client without undue delay, and in any case within forty-eight (48) hours, of any personal data breach of which it becomes aware; and shall assist the Client in responding to requests of data subjects.")
    b.clause("Safeguarding, counseling, and health records shall never appear in parent or student views or exports and shall be accessible only to the named roles agreed in the access list.")
    b.clause("No Client Data shall be used by the Service Provider for demonstrations, testing for other clients, or marketing. Demonstrations use fictitious data.")

    # 14 Publicity & recognition
    b.article("Publicity, Reference, and Recognition")
    b.clause("The Client hereby authorises the Service Provider to use the name and logo of Lebawi International Academy as a reference client in the Service Provider's promotional materials, website, social media, proposals, tender submissions, and case studies, and to state that the Client uses MaJo e-School Bridge under the Hybrid Implementation. The Service Provider shall not disclose Client Data, confidential information, the Contract Price (except where required in a tender), or any statement attributed to a named officer of the Client without that officer's written approval.")
    b.clause("The Client agrees to issue to the Service Provider, within thirty (30) days of Go-Live, a Certificate of Recognition substantially in the form of Annex I, signed by the Principal or the General Manager and bearing the Client's seal.")
    b.clause("The Client further agrees, upon reasonable written request of the Service Provider, to issue a certificate of satisfactory completion, a letter of recommendation, a certificate of good performance for use in public or private tenders, or a similar certificate, truthfully reflecting the Service Provider's performance under this Agreement, at no charge, within ten (10) Working Days of the request.")
    b.clause("The Client may name the Service Provider as its technology partner in its own communications. Either Party may install a modest acknowledgement in the System's login screen or footer (“Powered by MaJo e-School Bridge”).")
    b.clause("The Service Provider may photograph or film training sessions and the launch for promotional use only with the prior written consent of the Client, and shall not publish images in which students are identifiable without the consent required by law.")

    # 15 Change control
    b.article("Changes and Variations")
    b.clause("The Requirements Register and the Customization Register signed under Milestone M2 are the baseline. Any change requested after that date is recorded in the issue log, assessed by the Service Provider for effect on scope, time, and price, and implemented only after a written variation order signed by the Client's project sponsor and the Service Provider's project lead.")
    b.clause("Changes that are configuration within the existing Modules, correct a defect, or are required for compliance with the Inception Report are made at no extra charge. New Modules, new integrations, new reports beyond the configured templates, or re-work caused by late or inaccurate Client inputs are priced in the variation order by reference to the unit prices in Annex A or at a daily rate to be agreed.")
    b.clause("Any variation that changes the Contract Price shall state the revised price, VAT, and payment timing and shall be annexed to this Agreement.")

    # 16 Taxes
    b.article("Taxes, Duties, and Stamp Duty")
    b.clause("The Contract Price includes VAT at fifteen percent (15%) as shown in Article 4. The Service Provider shall issue VAT invoices in the form prescribed by the Ethiopian tax authority for each instalment.")
    b.clause("Each Party bears the taxes imposed on it by law. Withholding tax is governed by Article 5.5.")
    b.clause("Stamp duty on this Agreement, if any, under the Stamp Duty Proclamation No. 110/1998 shall be borne equally by the Parties, and the Parties shall cooperate in registering or authenticating this Agreement before the competent authority if either Party so requests, the cost being borne by the requesting Party.")

    # 17 Force majeure
    b.article("Force Majeure")
    b.clause("Neither Party is liable for failure or delay in performing its obligations, other than payment obligations already due, to the extent caused by an event beyond its reasonable control that it could not have foreseen or prevented, including war, civil unrest, state of emergency, strike not limited to the affected Party's employees, epidemic, fire, flood, earthquake, prolonged failure of public power or telecommunications, or act of a public authority, within the meaning of Articles 1792 to 1794 of the Civil Code.")
    b.clause("The affected Party shall notify the other in writing within seven (7) days of the event, stating its expected duration, and shall use reasonable efforts to resume performance. Dates in Annex B are extended by the duration of the event.")
    b.clause("If a force majeure event continues for more than ninety (90) days, either Party may terminate this Agreement by thirty (30) days' written notice, and Article 18.5 applies.")

    # 18 Termination
    b.article("Suspension and Termination")
    b.clause("Either Party may terminate this Agreement by written notice if the other Party commits a material breach and fails to remedy it within thirty (30) days of receiving a written notice specifying the breach and requiring its remedy.")
    b.clause("Either Party may terminate this Agreement by written notice with immediate effect if the other Party becomes insolvent, is dissolved, or ceases to carry on business.")
    b.clause("The Client may terminate this Agreement for its convenience before Milestone M3 by thirty (30) days' written notice. In that case the first instalment is retained by the Service Provider as full compensation for mobilisation, the Cloud Node preparation, and the work performed to the date of termination, and no further amount is due by either Party.")
    b.clause("The Client may terminate for convenience after Milestone M3 by thirty (30) days' written notice, in which case it shall pay for all Deliverables accepted or deemed accepted and for work in progress on a pro-rata basis of the instalment then running, and the licence under Article 11 continues for the Modules delivered and paid.")
    b.clause("On any termination: the Service Provider shall hand over all Deliverables completed, the School Node remains with the Client together with all Client Data, the Service Provider shall assist in an orderly handover for up to ten (10) Working Days, and each Party shall return or destroy the other's confidential information. Articles 9.2, 9.5, 11, 13, 14.1, 19, 20, and 21 survive termination.")

    # 19 Liability
    b.article("Liability and Indemnity")
    b.clause("Each Party is liable to the other for direct damage caused by its breach of this Agreement in accordance with the Civil Code. Neither Party is liable for indirect or consequential loss, loss of profit, or loss of goodwill, except in case of fraud, gross negligence, wilful misconduct, breach of confidentiality, or breach of Article 13.")
    b.clause("The total aggregate liability of the Service Provider under or in connection with this Agreement shall not exceed the Contract Price actually paid, except in the cases excluded in Article 19.1.")
    b.clause("The Service Provider is not liable for loss of Client Data caused by the Client's failure to keep the backup disks, perform the weekly disk swap, protect the server administrator account or the encryption key, or maintain the server room, power, and network under Annex D.")
    b.clause("Each Party shall indemnify the other against third-party claims arising from its own breach of law, breach of confidentiality, or infringement of intellectual property rights.")

    # 20 Disputes
    b.article("Settlement of Disputes")
    b.clause("Any dispute arising out of or in connection with this Agreement shall first be referred to the project sponsor of the Client and the General Manager of the Service Provider, who shall meet within ten (10) Working Days of a written notice of dispute and attempt in good faith to settle it amicably within thirty (30) days.")
    b.clause("If the dispute is not settled amicably within that period, either Party may refer it to arbitration in Addis Ababa by a sole arbitrator appointed by agreement of the Parties or, failing agreement within fifteen (15) days, by the Addis Ababa Chamber of Commerce and Sectoral Associations Arbitration Institute, in accordance with the Arbitration and Conciliation Working Procedure Proclamation No. 1237/2021. The language of arbitration is English, with Amharic documents admitted without translation. The award is final and binding.")
    b.clause("Nothing in this Article prevents a Party from seeking interim or conservatory measures from the competent Federal Courts of Ethiopia, or from enforcing an undisputed payment obligation before those courts.")
    b.clause("Pending resolution of a dispute, the Parties shall continue to perform their obligations not affected by the dispute.")

    # 21 Governing law
    b.article("Governing Law and Language")
    b.clause("This Agreement is governed by and construed in accordance with the laws of the Federal Democratic Republic of Ethiopia, including the Civil Code of 1960 and the Commercial Code Proclamation No. 1243/2021.")
    b.clause("This Agreement is made in the English language. If an Amharic version is prepared for registration or any other purpose, the English version prevails in case of discrepancy, unless the Parties sign a bilingual version stating otherwise.")

    # 22 Notices
    b.article("Notices")
    b.clause("Notices under this Agreement shall be in writing and delivered by hand against receipt, by registered post, or by e-mail with confirmation of receipt, to the addresses on the first page or any other address notified in writing. Day-to-day project communications may be made in the project WhatsApp group or the issue log, but notices of breach, termination, dispute, or variation shall be given as provided in this Article.")
    b.clause("For the Client: attention of ______________________________, ______________________________, e-mail ______________________________. For the Service Provider: attention of the General Manager, Majo Bridge Technologies and Events PLC, e-mail majobridgetech@gmail.com, copy nabilmaya6464@gmail.com.")

    # 23 General
    b.article("General Provisions")
    b.clause("Entire agreement. This Agreement and its Annexes constitute the entire agreement between the Parties on its subject matter and supersede all prior proposals, negotiations, and correspondence, including the commercial terms of the Financial Proposal to the extent inconsistent with Articles 4 and 5.")
    b.clause("Amendment. No amendment or waiver is valid unless in writing and signed by the authorised representatives of both Parties.")
    b.clause("Assignment and subcontracting. Neither Party may assign this Agreement without the written consent of the other. The Service Provider may use subcontractors for hosting of the Cloud Node and for training logistics, remaining fully responsible for their performance and confidentiality.")
    b.clause("Independent contractors. The Parties are independent contractors. Nothing in this Agreement creates a partnership, joint venture, agency, or employment relationship. The Service Provider's personnel remain its employees for all purposes.")
    b.clause("Non-solicitation. During the term and for twelve (12) months after, neither Party shall directly solicit for employment the staff of the other Party who worked on the project, without the other's written consent.")
    b.clause("Severability. If any provision is held invalid or unenforceable, the remaining provisions remain in force and the Parties shall replace the invalid provision with a valid one of similar economic effect.")
    b.clause("Anti-corruption. Each Party confirms that no payment, gift, or advantage has been or will be offered or given to any officer, employee, or representative of the other Party to obtain or retain this Agreement, and that it complies with the Corruption Crimes Proclamation No. 881/2015.")
    b.clause("Counterparts and copies. This Agreement is signed in two (2) original counterparts, each of which is an original, one retained by each Party. Signed scanned copies exchanged by e-mail are valid evidence of signature.")

    # 24 Representations
    b.article("Authority of Signatories")
    b.clause("Each signatory declares that he or she is duly authorised to sign this Agreement on behalf of the Party he or she represents, that the Party has the legal capacity to enter into it, and that all internal approvals required have been obtained. Copies of the authorising documents (board resolution, power of attorney, or appointment letter) may be attached at either Party's request.")


def signatures(b: Builder) -> None:
    b.h2("IN WITNESS WHEREOF")
    b.p("the Parties have caused this Agreement to be signed by their duly authorised representatives on the day, month, and year first above written, in Addis Ababa, Ethiopia.")
    t = b.doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_widths(t, [8.3, 8.3])
    left, right = t.rows[0].cells
    for cell, who in ((left, "For and on behalf of\nLEBAWI INTERNATIONAL\nACADEMY"), (right, "For and on behalf of\nMAJO BRIDGE TECHNOLOGIES\nAND EVENTS PLC")):
        cell.text = ""
        r = cell.paragraphs[0].add_run(who)
        r.bold = True
        for _ in range(3):
            cell.add_paragraph("")
        for line in ("Signature: ______________________________", "Name: ______________________________", "Title: ______________________________", "Date: ______________________________", "Seal:"):
            cell.add_paragraph(line)
    b.p()
    b.h2("WITNESSES")
    b.p("Signed in the presence of the following witnesses, who confirm that the representatives signed before them:")
    b.table(
        ["No.", "Full name", "ID / Address", "Signature"],
        [["1", "\n", "\n", "\n"], ["2", "\n", "\n", "\n"], ["3", "\n", "\n", "\n"]],
        widths=[1.0, 6.0, 6.0, 3.6],
        size=10,
    )
    b.p("Each page of this Agreement and of its Annexes is initialled by both Parties.", italic=True, color=GREY, size=9)


def annexes(b: Builder) -> None:
    # Annex A
    b.page_break()
    b.h1("ANNEX A — MODULE LIST, CATEGORIES, AND UNIT PRICES")
    b.p("Forty-nine (49) Modules of MaJo e-School Bridge licensed under this Agreement at the published unit prices of Financial Proposal MBT-LIA-FIN-2026-02 (VAT exclusive). The four login portals are included at no module charge. The “Confirmed” column is ticked at the kick-off meeting in accordance with Article 3.3.")
    rows: list[list[str]] = []
    for group, items in MODULES:
        rows.append([group.upper(), "", "", "", ""])
        for no, name, cat, price in items:
            rows.append([str(no), name, cat, f"{price:,}", "☐"])
    rows.append(["", "Total, 49 Modules", "", f"{MODULES_ETB:,}", ""])
    t = b.table(["No.", "Module", "Category", "Unit price (ETB)", "Confirmed"], rows, widths=[1.2, 7.8, 2.6, 3.0, 2.2], size=9, bold_last=True)
    for row in t.rows[1:]:
        if row.cells[0].text and row.cells[1].text == "" and row.cells[3].text == "":
            merged = row.cells[0].merge(row.cells[4])
            for par in list(merged.paragraphs)[1:]:
                par._element.getparent().remove(par._element)
            for par in merged.paragraphs:
                for r in par.runs:
                    r.bold = True
                    r.font.color.rgb = BLUE
            set_cell_shading(merged, "F4F7FB")
    b.p("Price categories: Premium 95,000 · High 75,000 · Mid 65,000 · Standard 50,000 · Lite 35,000 ETB per Module, VAT exclusive.", size=9, color=GREY)
    b.table(
        ["Included login portals", "Users"],
        [
            ["P1 Administration / ERP", "Board, General Manager, Principal, Vice Principal, Directors, Registrar, DoSA, QA, IT, administration staff"],
            ["P2 Teacher", "Homeroom and subject teachers: attendance, homework, marks, lesson plans, materials, parent messages"],
            ["P3 Parent", "Linked guardians: published grades, attendance, homework, announcements, messages"],
            ["P4 Student", "Enrolled students: timetable, homework, materials, examinations, announcements"],
        ],
        widths=[5.0, 11.6],
        size=9.5,
    )

    # Annex B
    b.page_break()
    b.h1("ANNEX B — DELIVERABLES AND MILESTONES")
    b.p("As set out in Sections 7 and 8 of the Inception Report. Acceptance follows Article 10.")
    b.table(["D", "Deliverable", "Content", "Due", "Accepted by"], [list(d) for d in DELIVERABLES], widths=[1.0, 4.0, 7.0, 1.4, 3.2], size=9)
    b.table(["No.", "Milestone", "Target date"], [list(m) for m in MILESTONES], widths=[1.2, 11.0, 4.4], size=9.5)
    b.p("A one-page progress report (done / next / issues / risks / decisions needed) is issued every second Friday and noted at the steering meeting. Milestone M3 and later depend on the Client's server being ready by Friday 6 November 2026 (Article 6.3).", size=9, color=GREY)

    # Annex C
    b.h1("ANNEX C — PAYMENT SCHEDULE")
    b.table(
        ["Instalment", "Share", "Amount (ETB)", "In words (Ethiopian Birr)", "Trigger"],
        [
            ["First", "40%", f"{PAY1:,}", WORDS[PAY1], "Signing of this Agreement"],
            ["Second", "40%", f"{PAY2:,}", WORDS[PAY2], "UAT readiness notice after Milestone M5"],
            ["Third", "20%", f"{PAY3:,}", WORDS[PAY3], "Go-Live Certificate, Milestone M7"],
            ["Total", "100%", f"{TOTAL_ETB:,}", WORDS[TOTAL_ETB], "VAT inclusive"],
        ],
        widths=[1.8, 1.4, 2.6, 6.6, 4.2],
        size=9,
        bold_last=True,
    )
    b.p(f"Breakdown: Modules {MODULES_ETB:,} + implementation {IMPLEMENTATION_ETB:,} − discount {DISCOUNT_ETB:,} = {SUBTOTAL_ETB:,} ETB VAT exclusive; VAT 15% = {VAT_ETB:,} ETB; total {TOTAL_ETB:,} ETB. Withholding tax is deducted from each payment as provided in Article 5.5 and evidenced by the official receipt.", size=9, color=GREY)

    # Annex D
    b.h1("ANNEX D — INFRASTRUCTURE, PREMISES, AND INPUTS PROVIDED BY THE CLIENT")
    b.p("Procured, installed, and maintained by the Client at its own cost (Article 8.1(b)). The Service Provider installs the School Node software on this server with the Client's IT Administrator present. Full detail in Annex D of the Inception Report.")
    b.table(["Item", "Minimum requirement"], [list(s) for s in SERVER_SPEC], widths=[3.4, 13.2], size=9.5)
    b.table(
        ["Client input", "Due date", "Needed for"],
        [
            ["Named project sponsor, focal person, IT Administrator, and office counterparts", "Tue 27 Oct 2026 (kick-off)", "Governance; all Deliverables"],
            ["Officers and teachers available for requirements workshops", "Tue 27 – Fri 30 Oct 2026", "D2 registers"],
            ["Physical server, UPS, room, LAN, internet line, sub-domain, two backup disks", "Fri 6 Nov 2026", "D3 School Node; M3 and later"],
            ["Completed data templates (students, parents, teachers, staff, classes, timetable, grading scheme)", "Fri 13 Nov 2026", "D5 migration; M4"],
            ["Sign-off of the imported roster and staff list", "Fri 20 Nov 2026", "M4"],
            ["Staff availability for training per Annex F; rooms and devices", "W3–W6 (9 Nov – 4 Dec 2026)", "D6; M5"],
            ["Officers to run UAT scenarios", "W8 (14–18 Dec 2026)", "D9; M6"],
            ["Parent briefings scheduled per section", "W6–W9", "Parent linking; Go-Live"],
        ],
        widths=[8.2, 4.2, 4.2],
        size=9,
    )

    # Annex E–H
    b.h1("ANNEX E — INCEPTION REPORT")
    b.p("Document MBT-LIA-INC-2026-01, “Inception Report — Student Management System (SMS) Project at Lebawi International Academy”, Version 1.0 of 27 October 2026, fifteen (15) pages, including its Annexes A (module scope checklist), B (data templates), C (training timetable), D (hybrid architecture and server specification), and E (requirements and customization registers, attached when signed under Milestone M2). The final version accepted under Milestone M2 replaces Version 1.0 for the purposes of this Agreement. Attached hereto and initialled on every page by both Parties.")
    b.h1("ANNEX F — TRAINING MODULES AND TRAINER'S GUIDE")
    b.p("Document MBT-LIA-TRN-2026-01, Version 1.0 of 27 October 2026, seventeen (17) pages: thirteen (13) training modules T1 to T13 with attendance sheet, evaluation form, competency checklists, and room and practice-data requirements. Attached hereto.")
    b.h1("ANNEX G — FINANCIAL PROPOSAL")
    b.p("Document MBT-LIA-FIN-2026-02, “Financial Proposal — One-time module licence and implementation”, of 16 October 2026, five (5) pages, Schedules A and B. Attached for the record of unit prices and the recalculation rule in Article 3.3. Where it conflicts with Articles 4 and 5 of this Agreement, this Agreement prevails.")
    b.h1("ANNEX H — NOTICE OF INTENT TO AWARD")
    b.p("Letter of the Procurement Department of Lebawi International Academy, “Notice of Intent to Award — Student Management System (SMS) Project”, reference ______________, dated ______________ 2026, addressed to Majo Bridge Technologies and Events PLC. Attached hereto.")

    # Annex I
    b.page_break()
    b.h1("ANNEX I — FORM OF CERTIFICATE OF RECOGNITION")
    b.p("(To be issued on the letterhead of Lebawi International Academy within thirty (30) days of Go-Live, Article 14.2)", italic=True, color=GREY, size=9)
    b.p()
    b.p("CERTIFICATE OF RECOGNITION", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=16, color=BLUE)
    b.p("Lebawi International Academy", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    b.p("hereby recognises", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    b.p("MAJO BRIDGE TECHNOLOGIES AND EVENTS PLC", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
    b.p("for the successful design, customization, and implementation of the Student Management System “MaJo e-School Bridge” at Lebawi International Academy under a hybrid architecture that keeps the Academy's student records on its own premises, and for the professional training of the Academy's leadership, officers, teachers, parents, and students, delivered in accordance with the Agreement MBT-LIA-AGR-2026-01 and brought into productive use for all sections on ______________ 2026.", align=WD_ALIGN_PARAGRAPH.CENTER)
    b.p("The Academy commends the Service Provider's commitment, technical competence, and respect for the Academy's data and values, and recommends Majo Bridge Technologies and Events PLC to other educational institutions.", align=WD_ALIGN_PARAGRAPH.CENTER)
    b.p()
    t = b.doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_widths(t, [8.3, 8.3])
    for cell, who in zip(t.rows[0].cells, ("______________________________\nPrincipal\nLebawi International Academy", "______________________________\nGeneral Manager\nLebawi International Academy")):
        cell.text = ""
        cell.paragraphs[0].add_run(who)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    b.p()
    b.p("Addis Ababa, ______________ 2026 G.C. (______________ 2019 E.C.)        Seal of the Academy", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)


def main() -> None:
    b = Builder()
    cover(b)
    parties(b)
    articles(b)
    signatures(b)
    annexes(b)
    b.save()
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
